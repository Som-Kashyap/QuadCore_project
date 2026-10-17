"""
Groq LLM Client for automated document extraction (.pdf and .docx).
Encapsulates document text extraction, Groq inference, and structured JSON parsing.
"""

import json
import logging
from pathlib import Path
from typing import Any, Dict

import docx
from groq import AsyncGroq, APIError, AuthenticationError, RateLimitError
import pypdf

from app.core.config import get_settings
from app.domains.documents.prompts import BASE_SYSTEM_PROMPT, get_prompt_for_category
from app.domains.documents.schemas import DocumentCategory

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Custom Domain Exceptions
# ---------------------------------------------------------------------------
class GroqClientError(Exception):
    """Base exception for all Groq extraction client errors."""
    pass


class GroqAPIKeyMissingError(GroqClientError):
    """Raised when GROQ_API_KEY is not configured in settings or environment."""
    pass


class UnsupportedDocumentFormatError(GroqClientError):
    """Raised when an unsupported document format (anything other than .pdf or .docx) is provided."""
    pass


class DocumentReadError(GroqClientError):
    """Raised when a local file cannot be read or parsed by pypdf/docx."""
    pass


class GroqAPIError(GroqClientError):
    """Raised when the Groq API call fails or times out."""
    pass


class ExtractionParsingError(GroqClientError):
    """Raised when the LLM response cannot be parsed into the expected JSON structure."""
    pass


# ---------------------------------------------------------------------------
# Groq Extraction Client
# ---------------------------------------------------------------------------
class GroqExtractionClient:
    """
    Dedicated client for extracting structured decarbonization metrics from .pdf and .docx
    using Groq's high-speed inference.
    """

    DEFAULT_MODEL = "llama-3.3-70b-versatile"

    def __init__(self, api_key: str | None = None):
        settings = get_settings()
        self.api_key = api_key or settings.GROQ_API_KEY

    def _get_client(self) -> AsyncGroq:
        if not self.api_key:
            logger.error("Groq extraction requested but GROQ_API_KEY is not configured.")
            raise GroqAPIKeyMissingError(
                "GROQ_API_KEY is not configured in environment or .env file."
            )
        return AsyncGroq(api_key=self.api_key)

    # -----------------------------------------------------------------------
    # Document Readers
    # -----------------------------------------------------------------------
    def read_pdf(self, file_path: Path | str, max_pages: int = 20) -> str:
        """Extract text from a .pdf file up to max_pages."""
        path = Path(file_path)
        if not path.exists():
            raise DocumentReadError(f"PDF file does not exist: {path}")

        try:
            reader = pypdf.PdfReader(str(path))
            pages_text = []
            total_pages = min(len(reader.pages), max_pages)

            for i in range(total_pages):
                page = reader.pages[i]
                extracted = page.extract_text()
                if extracted and extracted.strip():
                    pages_text.append(f"--- PAGE {i + 1} ---\n{extracted.strip()}")

            full_text = "\n\n".join(pages_text).strip()
            if not full_text:
                raise DocumentReadError("No readable text found in PDF (may be scanned image).")

            return full_text
        except DocumentReadError:
            raise
        except Exception as exc:
            logger.error("Failed to read PDF file %s: %s", path, exc, exc_info=True)
            raise DocumentReadError(f"Error reading PDF file: {exc}") from exc

    def read_docx(self, file_path: Path | str) -> str:
        """Extract paragraphs and table text from a .docx file."""
        path = Path(file_path)
        if not path.exists():
            raise DocumentReadError(f"DOCX file does not exist: {path}")

        try:
            doc = docx.Document(str(path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            table_lines = []
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))

            full_text = "\n".join(paragraphs + table_lines).strip()
            if not full_text:
                raise DocumentReadError("No text found in DOCX file.")

            return full_text
        except DocumentReadError:
            raise
        except Exception as exc:
            logger.error("Failed to read DOCX file %s: %s", path, exc, exc_info=True)
            raise DocumentReadError(f"Error reading DOCX file: {exc}") from exc

    # -----------------------------------------------------------------------
    # Groq Inference Execution
    # -----------------------------------------------------------------------
    async def _execute_extraction(
        self,
        text_content: str,
        category: DocumentCategory,
        filename: str,
    ) -> Dict[str, Any]:
        """Send extracted document text to Groq LLM with category instructions."""
        client = self._get_client()
        category_prompt = get_prompt_for_category(category)

        # Truncate text content safely to fit well within context limits
        truncated_text = text_content[:40000]

        user_content = (
            f"Document Filename: {filename}\n"
            f"Document Category: {category.value}\n\n"
            f"Extraction Target Requirements:\n{category_prompt}\n\n"
            f"Document Extracted Text:\n{truncated_text}"
        )

        try:
            response = await client.chat.completions.create(
                model=self.DEFAULT_MODEL,
                messages=[
                    {"role": "system", "content": BASE_SYSTEM_PROMPT},
                    {"role": "user", "content": user_content},
                ],
                temperature=0.0,
                response_format={"type": "json_object"},
            )

            raw_choice = response.choices[0].message.content
            if not raw_choice:
                raise ExtractionParsingError("Groq returned an empty response.")

            parsed_data = json.loads(raw_choice)
            return parsed_data

        except AuthenticationError as exc:
            logger.error("Groq authentication failure: %s", exc)
            raise GroqAPIKeyMissingError("Invalid or missing Groq API key.") from exc
        except RateLimitError as exc:
            logger.warning("Groq rate limit encountered: %s", exc)
            raise GroqAPIError("Groq rate limit exceeded. Please retry shortly.") from exc
        except APIError as exc:
            logger.error("Groq API error during extraction: %s", exc)
            raise GroqAPIError(f"Groq API error: {exc}") from exc
        except json.JSONDecodeError as exc:
            logger.error("Failed to parse Groq response as JSON: %s", exc)
            raise ExtractionParsingError(f"Malformed JSON returned by model: {exc}") from exc
        except Exception as exc:
            logger.error("Unexpected error during Groq extraction: %s", exc, exc_info=True)
            raise GroqClientError(f"Unexpected extraction failure: {exc}") from exc

    # -----------------------------------------------------------------------
    # Public Extraction Endpoints
    # -----------------------------------------------------------------------
    async def extract_from_pdf(
        self, file_path: Path | str, category: DocumentCategory
    ) -> Dict[str, Any]:
        """Extract structured fields from a PDF file."""
        text = self.read_pdf(file_path)
        filename = Path(file_path).name
        return await self._execute_extraction(text, category, filename)

    async def extract_from_docx(
        self, file_path: Path | str, category: DocumentCategory
    ) -> Dict[str, Any]:
        """Extract structured fields from a DOCX file."""
        text = self.read_docx(file_path)
        filename = Path(file_path).name
        return await self._execute_extraction(text, category, filename)

    async def extract(
        self,
        file_path: Path | str,
        extension: str,
        category: DocumentCategory,
    ) -> Dict[str, Any]:
        """
        Unified router method. Validates that the file extension is .pdf or .docx,
        dispatches to the appropriate parser, and returns the extracted fields.
        """
        ext = extension.lower().lstrip(".")
        if ext == "pdf":
            return await self.extract_from_pdf(file_path, category)
        elif ext == "docx":
            return await self.extract_from_docx(file_path, category)
        else:
            raise UnsupportedDocumentFormatError(
                f"Unsupported format '{ext}'. Automated extraction currently only supports .pdf and .docx."
            )
